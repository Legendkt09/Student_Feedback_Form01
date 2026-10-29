
import os
import re
import hmac
import secrets
from datetime import datetime
from functools import wraps

from flask import (
    Flask, request, session, redirect, url_for,
    render_template_string, flash, send_file, abort
)
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

app = Flask(__name__)
app.secret_key = os.environ.get("FLASK_SECRET_KEY", secrets.token_hex(32))

BASE = os.path.dirname(os.path.abspath(__file__))
STATIC = os.path.join(BASE, "static")
DATA = os.path.join(BASE, "data")
os.makedirs(STATIC, exist_ok=True)
os.makedirs(DATA, exist_ok=True)

COLLEGE_NAME = "Dr.A.D. Shinde College of Engineering "
COLLEGE_SUBTITLE = "Bhadgaon"

# Demo credentials: change before real deployment.
ADMIN_USER = os.environ.get(
    "FEEDBACK_ADMIN_USER",
    "snehal.jirange@adshindecoe.edu.in"
)

ADMIN_PASS = os.environ.get(
    "FEEDBACK_ADMIN_PASSWORD",
    "ADSC_Admin@123"
)
STUDENT_USER = os.environ.get("FEEDBACK_STUDENT_USER", "student")
STUDENT_PASS = os.environ.get("FEEDBACK_STUDENT_PASSWORD", "Student@123")

ACADEMIC_YEARS = [
    "2026-2027", "2027-2028", "2028-2029",
    "2029-2030", "2030-2031", "2031-2032"
]

STUDY_YEARS = ["First Year", "Second Year", "Third Year", "Final Year"]
SEMESTERS = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII"]

# Check these lists against your college's official lists.
DEPARTMENTS = [
    "Artificial Intellegence And Mechine learning",
    "Computer Science and Engineering",
    "Mechanical and Mechatronics Engineering",
    "Civil Engineering",
    "Electrical and Computer Engineering",
    "Electronics and Computer Science Engineering",
]

SUBJECTS = [
    "Power Electronics",
    "Electric Transmission and Distribution",
    "Microcontroller 8051",
    "Fundamental of Design for Product Manufacturing",
    "Innovation, Business Model and Entrepreneurship",
    "Elecrical Power Generation and Utilization",
]

FACULTY_NAMES = [
    "Prof.A.R.Bandekar",
    "Prof.D.S.Mahajan",
    "Prof.P.S.Lohar",
    "Prof.M.N.Patil",
    "Prof.S.N.jirange",
    "Prof.A.B.Angadi",
]

# IMPORTANT: Replace the wording below with your exact original five options.
OPTIONS = ["Excellent", "Very Good", "Good", "Average", "Poor"]

QUESTIONS = [
    "Faculty knowledge of the course/subject",
    "Clarity of concepts and explanations",
    "Encouragement of questions and classroom discussions",
    "Use of examples, practical applications and real-time/industry examples",
    "Coverage of the prescribed syllabus",
    "Faculty availability and approachability to students",
    "Teaching methodology and use of appropriate teaching-learning resources",
]

FIELDS = [
    ("academic_year", "Academic Year", ACADEMIC_YEARS),
    ("prn", "Student PRN", None),
    ("roll_number", "Roll Number", None),
    ("study_year", "Study Year", STUDY_YEARS),
    ("semester", "Semester", SEMESTERS),
    ("department", "Department", DEPARTMENTS),
    ("subject", "Subject", SUBJECTS),
    ("faculty", "Faculty Name", FACULTY_NAMES),
]

HEADERS = [
    "Submitted At", "Academic Year", "PRN", "Roll Number",
    "Study Year", "Semester", "Department", "Subject", "Faculty Name"
] + QUESTIONS


# -------------------- DESIGN --------------------

CSS = r"""
:root{--ink:#16264b;--blue:#315eff;--purple:#8048ef;--pale:#edf2ff}
*{box-sizing:border-box}
body{margin:0;min-height:100vh;font-family:Segoe UI,Arial,sans-serif;color:var(--ink);
background:linear-gradient(120deg,rgba(8,20,49,.78),rgba(44,54,115,.56)),url('/static/college_bg.jpg') center/cover fixed}
header{display:flex;align-items:center;gap:14px;padding:14px 5%;color:white;background:linear-gradient(100deg,#101b39,#294d8b);box-shadow:0 12px 30px #09132d66}
.logo{width:58px;height:58px;object-fit:contain;background:white;border-radius:14px;padding:5px}
.brand{display:flex;flex-direction:column;gap:5px}
.brand strong{font-size:17px}.brand small{opacity:.85}
nav{margin-left:auto;display:flex;gap:14px;flex-wrap:wrap}
nav a{color:white;text-decoration:none;font-weight:700}
main{width:min(1000px,92%);margin:42px auto;min-height:65vh}
.card{background:rgba(255,255,255,.94);border:1px solid #ffffffc9;border-radius:27px;padding:clamp(24px,5vw,52px);
box-shadow:0 28px 70px #081a4566,inset 0 1px white;backdrop-filter:blur(16px);animation:enter .35s ease-out}
@keyframes enter{from{opacity:0;transform:translateY(15px) perspective(900px) rotateX(2deg)}to{opacity:1;transform:translateY(0) perspective(900px) rotateX(0)}}
.hero{display:grid;grid-template-columns:1.2fr .8fr;gap:30px;align-items:center;color:white;background:linear-gradient(130deg,#111d40,#31569b 65%,#684cc1)}
.hero p{color:#e4eaff}.hero h1{font-size:clamp(36px,5vw,60px);line-height:1.08}
.panel{border:1px solid #ffffff66;border-radius:23px;padding:25px;background:#ffffff17;box-shadow:0 20px 35px #0a173c55}
.hero-logo{width:110px;height:110px;object-fit:contain;background:white;border-radius:18px;padding:7px}
.pill{display:inline-block;background:#e5ecff;color:#2c4d9c;border-radius:30px;padding:9px 14px;font-size:12px;font-weight:800;letter-spacing:.6px}
h1{font-size:clamp(29px,4vw,43px);margin:20px 0 12px;line-height:1.18}
h2{line-height:1.45}p{line-height:1.7;color:#6c7b9e}
label{display:block;font-weight:750;margin:22px 0 9px}
input,select{width:100%;min-height:54px;padding:14px;border:1px solid #d2dcf5;border-radius:13px;background:#f8faff;font-size:16px;color:#18274b;outline:none}
input:focus,select:focus{border-color:#536fff;box-shadow:0 0 0 4px #536fff20}
.actions{display:flex;flex-wrap:wrap;gap:12px;margin-top:27px}
.btn{display:inline-flex;align-items:center;justify-content:center;padding:14px 20px;border:0;border-radius:13px;text-decoration:none;color:white;font-size:15px;font-weight:750;cursor:pointer;background:linear-gradient(120deg,var(--blue),var(--purple));box-shadow:0 9px 20px #4b56df44,inset 0 1px #ffffff55;transition:.2s}
.btn:hover{transform:translateY(-3px);box-shadow:0 14px 27px #4b56df55}
.light{background:linear-gradient(130deg,#fff,#e4eaff);color:#25427c}
.muted{background:#e9eef9;color:#405478}
.full{width:100%;margin-top:25px}
.progress{height:9px;border-radius:20px;background:#dfe7f9;overflow:hidden;margin:24px 0}
.progress span{display:block;height:100%;background:linear-gradient(90deg,#315eff,#48d4ee,#8048ef)}
.question{padding:23px;border-radius:20px;background:linear-gradient(135deg,#f3f6ff,#e3eaff);border:1px solid #dbe4ff;box-shadow:0 12px 25px #243d7718}
.alert{padding:13px 16px;margin:15px 0;border-radius:12px;background:#ffe6e6;color:#9c2020}
.success{text-align:center}.check{margin:auto;width:92px;height:92px;border-radius:50%;display:grid;place-items:center;color:white;font-size:48px;background:linear-gradient(135deg,#13bb82,#2f74ff);box-shadow:0 15px 35px #277ee855}
.table-wrap{overflow:auto;margin-top:24px}table{border-collapse:collapse;background:white;width:100%;font-size:13px}th,td{padding:12px;border-bottom:1px solid #e5eaf5;text-align:left;white-space:nowrap}th{background:#20396e;color:white}
footer{text-align:center;background:#101d40;color:#e5eaff;padding:22px;margin-top:35px}
@media(max-width:700px){.hero{grid-template-columns:1fr}header{padding:12px 4%}.brand strong{font-size:13px}main{margin:24px auto}.card{border-radius:20px}}
"""


# All routes render a complete independent HTML page.
def page(title, content, **context):
    template = r"""
    <!doctype html><html lang="en"><head>
    <meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
    <title>{{ title }}</title><style>{{ css|safe }}</style></head><body>
    <header>
      <img class="logo" src="/static/college_logo.png" alt="College logo"
           onerror="this.style.display='none'">
      <div class="brand"><strong>{{ college }}</strong><small>{{ subtitle }}</small></div>
      <nav><a href="/">Home</a>{% if role %}<a href="/logout">Logout</a>{% endif %}</nav>
    </header>
    <main>
      {% for message in messages %}<div class="alert">{{ message }}</div>{% endfor %}
      """ + content + r"""
    </main><footer>Student Faculty Feedback System</footer></body></html>
    """
    return render_template_string(
        template, title=title, css=CSS, college=COLLEGE_NAME,
        subtitle=COLLEGE_SUBTITLE, role=session.get("role"),
        messages=[m for m in []], **context
    )


def login_required(role):
    def decorator(fn):
        @wraps(fn)
        def wrapper(*args, **kwargs):
            if session.get("role") != role:
                return redirect(url_for("login", role=role))
            return fn(*args, **kwargs)
        return wrapper
    return decorator


def workbook_path(year):
    safe = re.sub(r"[^0-9-]", "", year)
    return os.path.join(DATA, f"feedback_{safe}.xlsx")


def prepare_sheet(ws):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="243B75")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 24


def save_feedback(data):
    year = data["academic_year"]
    path = workbook_path(year)

    if os.path.exists(path):
        wb = load_workbook(path)
        ws = wb.active
        if [c.value for c in ws[1]] != HEADERS:
            wb.close()
            raise ValueError(
                "Excel headers changed. Please back up and rename the old "
                "file data/feedback_" + year + ".xlsx, then submit again."
            )
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Feedback"
        ws.append(HEADERS)

    ws.append([
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        data["academic_year"], data["prn"], data["roll_number"],
        data["study_year"], data["semester"], data["department"],
        data["subject"], data["faculty"], *data["answers"]
    ])
    prepare_sheet(ws)
    wb.save(path)
    wb.close()


# -------------------- HOME / LOGIN --------------------

@app.route("/")
def home():
    content = r"""
    <section class="card hero">
      <div><span class="pill">ACADEMIC FEEDBACK PORTAL</span>
      <h1>Better Teaching.<br>Better Learning.</h1>
      <p>Secure, step-by-step faculty feedback portal.</p>
      <div class="actions">
        <a class="btn" href="/login/student">Student Login →</a>
        <a class="btn light" href="/login/admin">Admin Login →</a>
      </div></div>
      <div class="panel">
        <img class="hero-logo" src="/static/college_logo.png" alt="College logo"
             onerror="this.style.display='none'">
        <h2>""" + COLLEGE_NAME + r"""</h2><p>""" + COLLEGE_SUBTITLE + r"""</p>
        <p>Student Faculty Feedback System</p>
      </div>
    </section>"""
    return page("Home", content)


@app.route("/login/<role>", methods=["GET", "POST"])
def login(role):
    if role not in ("student", "admin"):
        abort(404)

    error = ""
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")

        if role == "admin":
            valid = hmac.compare_digest(username, ADMIN_USER) and hmac.compare_digest(password, ADMIN_PASS)
        else:
            valid = hmac.compare_digest(username, STUDENT_USER) and hmac.compare_digest(password, STUDENT_PASS)

        if valid:
            session.clear()
            session["role"] = role
            session["username"] = username
            if role == "admin":
                return redirect(url_for("admin"))
            session["feedback"] = {"answers": [""] * len(QUESTIONS)}
            return redirect(url_for("detail", step=0))
        error = "Username or password is incorrect."

    content = r"""
    <section class="card" style="max-width:560px;margin:auto">
      <span class="pill">""" + role.upper() + r""" PORTAL</span>
      <h1>""" + role.title() + r""" Login</h1>
      <p>Enter your login credentials to continue.</p>
      {% if error %}<div class="alert">{{ error }}</div>{% endif %}
      <form method="post">
        <label>Username</label><input name="username" autocomplete="username" required>
        <label>Password</label><input type="password" name="password" autocomplete="current-password" required>
        <button class="btn full" type="submit">Login Securely →</button>
      </form>
    </section>"""
    return page("Login", content, error=error)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("home"))


# -------------------- EIGHT SEPARATE DETAIL PAGES --------------------

@app.route("/feedback/details/<int:step>", methods=["GET", "POST"])
@login_required("student")
def detail(step):
    if not 0 <= step < len(FIELDS):
        abort(404)

    data = session.setdefault("feedback", {"answers": [""] * len(QUESTIONS)})
    key, label, choices = FIELDS[step]
    error = ""

    if request.method == "POST":
        value = request.form.get(key, "").strip()

        if not value:
            error = "Please fill in or choose an option."
        elif key == "prn" and not re.fullmatch(r"\d{10}", value):
            error = "PRN must contain exactly 10 digits."
        elif choices is not None and value not in choices:
            error = "Please select a valid dropdown option."
        else:
            data[key] = value
            session["feedback"] = data
            session.modified = True
            if step < len(FIELDS) - 1:
                return redirect(url_for("detail", step=step + 1))
            return redirect(url_for("question", number=0))

    value = data.get(key, "")
    options_html = ""
    if choices is not None:
        options_html = '<option value="">Choose ' + label + '</option>'
        for option in choices:
            selected = " selected" if option == value else ""
            options_html += '<option value="' + option + '"' + selected + '>' + option + '</option>'
        input_html = '<select name="' + key + '" required>' + options_html + '</select>'
    elif key == "prn":
        input_html = '<input name="prn" value="' + value + '" inputmode="numeric" pattern="[0-9]{10}" minlength="10" maxlength="10" placeholder="Exactly 10 digits" required>'
    else:
        input_html = '<input name="' + key + '" value="' + value + '" required>'

    previous = ""
    if step > 0:
        previous = '<a class="btn light" href="' + url_for("detail", step=step-1) + '">← Back</a>'

    content = r"""
    <section class="card">
      <span class="pill">DETAIL """ + str(step + 1) + " OF " + str(len(FIELDS)) + r"""</span>
      <h1>""" + label + r"""</h1>
      <p>This page contains only one detail. Complete it to continue.</p>
      <div class="progress"><span style="width:""" + str(int((step + 1) / (len(FIELDS) + len(QUESTIONS)) * 100)) + r"""%"></span></div>
      {% if error %}<div class="alert">{{ error }}</div>{% endif %}
      <form method="post"><label>""" + label + r"""</label>""" + input_html + r"""
        <div class="actions">""" + previous + r"""
          <button class="btn" type="submit">Next →</button>
          <a class="btn muted" href="/logout">Cancel</a>
        </div>
      </form>
    </section>"""
    return page(label, content, error=error)


# -------------------- ONE QUESTION PER PAGE --------------------

@app.route("/feedback/question/<int:number>", methods=["GET", "POST"])
@login_required("student")
def question(number):
    if not 0 <= number < len(QUESTIONS):
        abort(404)

    data = session.get("feedback", {})
    if any(not data.get(f[0]) for f in FIELDS):
        return redirect(url_for("detail", step=0))

    answers = data.get("answers", [""] * len(QUESTIONS))
    error = ""

    if request.method == "POST":
        answer = request.form.get("answer", "")
        if answer not in OPTIONS:
            error = "Please select one of the five options."
        else:
            answers[number] = answer
            data["answers"] = answers
            session["feedback"] = data
            session.modified = True
            if number < len(QUESTIONS) - 1:
                return redirect(url_for("question", number=number + 1))
            return redirect(url_for("review"))

    option_html = '<option value="">Choose one option</option>'
    for option in OPTIONS:
        selected = " selected" if answers[number] == option else ""
        option_html += '<option value="' + option + '"' + selected + '>' + option + '</option>'

    back_url = url_for("question", number=number - 1) if number > 0 else url_for("detail", step=len(FIELDS)-1)
    next_text = "Review Feedback →" if number == len(QUESTIONS) - 1 else "Next Question →"

    content = r"""
    <section class="card">
      <span class="pill">QUESTION """ + str(number + 1) + " OF " + str(len(QUESTIONS)) + r"""</span>
      <h1>Faculty Feedback</h1>
      <div class="progress"><span style="width:""" + str(int((number + 1) / len(QUESTIONS) * 100)) + r"""%"></span></div>
      <div class="question"><span class="pill">QUESTION """ + str(number + 1) + r"""</span>
      <h2>""" + QUESTIONS[number] + r"""</h2></div>
      {% if error %}<div class="alert">{{ error }}</div>{% endif %}
      <form method="post">
        <label>Choose your answer</label>
        <select name="answer" required>""" + option_html + r"""</select>
        <div class="actions">
          <a class="btn light" href=""" + '"' + back_url + '"' + r""">← Back</a>
          <button class="btn" type="submit">""" + next_text + r"""</button>
        </div>
      </form>
    </section>"""
    return page("Question " + str(number + 1), content, error=error)


# -------------------- REVIEW / SUBMIT --------------------

@app.route("/feedback/review", methods=["GET", "POST"])
@login_required("student")
def review():
    data = session.get("feedback", {})
    answers = data.get("answers", [])
    if any(not data.get(f[0]) for f in FIELDS):
        return redirect(url_for("detail", step=0))
    if len(answers) != len(QUESTIONS) or any(a not in OPTIONS for a in answers):
        return redirect(url_for("question", number=0))

    if request.method == "POST":
        try:
            save_feedback(data)
        except Exception as exc:
            app.logger.exception("Could not save feedback")
            return "Could not save feedback. Check the terminal and Excel file.", 500
        session.pop("feedback", None)
        session["submitted"] = True
        return redirect(url_for("success"))

    rows = ""
    for (key, label, _), _unused in zip(FIELDS, range(len(FIELDS))):
        rows += "<p><b>" + label + ":</b> " + str(data.get(key, "")) + "</p>"
    for i, q in enumerate(QUESTIONS):
        rows += "<p><b>" + str(i+1) + ". " + q + "</b><br>" + answers[i] + "</p>"

    content = r"""
    <section class="card"><span class="pill">FINAL REVIEW</span>
      <h1>Review Your Feedback</h1><p>Check your details and answers before submitting.</p>
      <div class="question">""" + rows + r"""</div>
      <form method="post"><div class="actions">
        <a class="btn light" href=""" + '"' + url_for("question", number=len(QUESTIONS)-1) + '"' + r""">← Back to Questions</a>
        <button class="btn" type="submit">Submit Feedback ✓</button>
      </div></form>
    </section>"""
    return page("Review Feedback", content)


@app.route("/feedback/success")
@login_required("student")
def success():
    if not session.pop("submitted", False):
        return redirect(url_for("home"))
    content = r"""
    <section class="card success">
      <div class="check">✓</div><span class="pill">SUBMISSION COMPLETE</span>
      <h1>Feedback Successfully Submitted!</h1>
      <p>Thank you for sharing your academic feedback.</p>
      <a class="btn" href="/logout">Finish →</a>
    </section>"""
    return page("Successfully Submitted", content)


# -------------------- ADMIN DASHBOARD --------------------

@app.route("/admin")
@login_required("admin")
def admin():
    year = request.args.get("year", "")
    if year and year not in ACADEMIC_YEARS:
        abort(400)

    records = []
    selected_years = [year] if year else ACADEMIC_YEARS

    for selected in selected_years:
        path = workbook_path(selected)
        if not os.path.exists(path):
            continue
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        for row_num in range(2, ws.max_row + 1):
            vals = [ws.cell(row_num, c).value for c in range(1, ws.max_column + 1)]
            if any(v is not None for v in vals):
                records.append((selected, row_num, vals))
        wb.close()

    year_options = '<option value="">All years</option>'
    for y in ACADEMIC_YEARS:
        s = " selected" if y == year else ""
        year_options += '<option' + s + ' value="' + y + '">' + y + '</option>'

    downloads = ""
    for y in ACADEMIC_YEARS:
        downloads += '<a class="btn light" href="' + url_for("download", year=y) + '">Download ' + y + ' Excel</a> '

    table_rows = ""
    for y, row_num, v in records:
        table_rows += "<tr><td>" + str(y) + "</td>"
        for idx in [0, 2, 3, 4, 5, 6, 7, 8]:
            table_rows += "<td>" + str(v[idx] or "") + "</td>"
        table_rows += '<td><a href="' + url_for("edit", year=y, row_num=row_num) + '">Edit</a></td></tr>'

    if not table_rows:
        table_rows = '<tr><td colspan="10">No feedback records found.</td></tr>'

    content = r"""
    <section class="card"><span class="pill">ADMINISTRATION</span>
      <h1>Feedback Dashboard</h1>
      <p>Only the administrator can view submitted feedback records.</p>
      <form method="get"><label>Filter by Academic Year</label>
      <select name="year">""" + year_options + r"""</select>
      <div class="actions"><button class="btn" type="submit">Filter</button></div></form>
      <h2>Year-wise Excel Downloads</h2><div class="actions">""" + downloads + r"""</div>
      <h2>Feedback Records</h2><div class="table-wrap"><table>
      <thead><tr><th>Year</th><th>Submitted At</th><th>PRN</th><th>Roll No.</th><th>Study Year</th><th>Semester</th><th>Department</th><th>Subject</th><th>Faculty</th><th>Edit</th></tr></thead>
      <tbody>""" + table_rows + r"""</tbody></table></div>
    </section>"""
    return page("Admin Dashboard", content)


@app.route("/admin/download/<year>")
@login_required("admin")
def download(year):
    if year not in ACADEMIC_YEARS:
        abort(404)
    path = workbook_path(year)
    if not os.path.exists(path):
        wb = Workbook()
        ws = wb.active
        ws.title = "Feedback"
        ws.append(HEADERS)
        prepare_sheet(ws)
        wb.save(path)
        wb.close()
    return send_file(path, as_attachment=True, download_name=f"feedback_{year}.xlsx")


@app.route("/admin/edit/<year>/<int:row_num>", methods=["GET", "POST"])
@login_required("admin")
def edit(year, row_num):
    if year not in ACADEMIC_YEARS:
        abort(404)
    path = workbook_path(year)
    if not os.path.exists(path):
        abort(404)

    wb = load_workbook(path)
    ws = wb.active
    if row_num < 2 or row_num > ws.max_row:
        wb.close()
        abort(404)

    values = [ws.cell(row_num, c).value for c in range(1, ws.max_column + 1)]

    if request.method == "POST":
        for col, header in enumerate(HEADERS, start=1):
            if header == "Submitted At":
                continue
            value = request.form.get("c" + str(col), "").strip()
            if header == "PRN" and not re.fullmatch(r"\d{10}", value):
                wb.close()
                return "PRN must contain exactly 10 digits.", 400
            if header in QUESTIONS and value not in OPTIONS:
                wb.close()
                return "Invalid feedback option.", 400
            ws.cell(row_num, col).value = value
        prepare_sheet(ws)
        wb.save(path)
        wb.close()
        return redirect(url_for("admin", year=year))

    fields_html = ""
    for col, header in enumerate(HEADERS, start=1):
        if header == "Submitted At":
            fields_html += "<label>" + header + "</label><input value='" + str(values[col-1] or "") + "' disabled>"
        elif header in QUESTIONS:
            opts = ""
            for op in OPTIONS:
                s = " selected" if values[col-1] == op else ""
                opts += "<option" + s + ">" + op + "</option>"
            fields_html += "<label>" + header + "</label><select name='c" + str(col) + "'>" + opts + "</select>"
        else:
            fields_html += "<label>" + header + "</label><input name='c" + str(col) + "' value='" + str(values[col-1] or "") + "' required>"

    wb.close()
    content = "<section class='card'><h1>Edit Feedback Record</h1><form method='post'>" + fields_html + "<div class='actions'><button class='btn' type='submit'>Save Changes</button><a class='btn light' href='/admin?year=" + year + "'>Cancel</a></div></form></section>"
    return page("Edit Feedback", content)


@app.errorhandler(404)
def not_found(error):
    return page("Page Not Found", "<section class='card'><h1>Page not found</h1><a class='btn' href='/'>Return Home</a></section>"), 404


if __name__ == "__main__":
    app.run(debug=True, port=5000)
